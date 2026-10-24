from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).parent
def shade(cell,fill):
 p=cell._tc.get_or_add_tcPr();x=OxmlElement('w:shd');x.set(qn('w:fill'),fill);p.append(x)
def borders(t):
 b=OxmlElement('w:tblBorders')
 for e in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+e);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
 t._tbl.tblPr.append(b)
def add_md(doc):
 lines=(ROOT/'docs/report.md').read_text(encoding='utf-8').splitlines();i=0
 while i<len(lines):
  s=lines[i].strip()
  if not s:i+=1;continue
  if s.startswith('!['):i+=1;continue
  if s.startswith('|'):
   rows=[]
   while i<len(lines) and lines[i].strip().startswith('|'):
    vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
    if not all(set(v)<=set('-: ') for v in vals):rows.append(vals)
    i+=1
   if rows:
    t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for ri,row in enumerate(rows):
     cells=t.add_row().cells
     for ci,v in enumerate(row):
      cells[ci].text=v;cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
      for p in cells[ci].paragraphs:
       p.paragraph_format.space_after=Pt(3)
       for r in p.runs:
        r.font.size=Pt(9);r.bold=ri==0
        if ri==0:r.font.color.rgb=RGBColor(255,255,255)
      shade(cells[ci],'17365D' if ri==0 else ('F2F6FA' if ri%2==0 else 'FFFFFF'))
    borders(t);doc.add_paragraph().paragraph_format.space_after=Pt(2)
   continue
  if s.startswith('# '):pass
  elif s.startswith('## '):
   heading=doc.add_heading(s[3:],2)
   if s[3:]=='Source control and group evidence':heading.paragraph_format.page_break_before=True
  elif s.startswith('### '):doc.add_heading(s[4:],3)
  elif s.startswith('#### '):doc.add_heading(s[5:],4)
  elif s.startswith('##### '):doc.add_heading(s[6:],5)
  elif s.startswith('- '):doc.add_paragraph(s[2:],style='List Bullet')
  elif s.startswith('```'):
   i+=1;code=[]
   while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
   p=doc.add_paragraph();r=p.add_run('\n'.join(code));r.font.name='Consolas';r.font.size=Pt(8)
  else:doc.add_paragraph(s.replace('**','').replace('`',''))
  i+=1
d=Document();sec=d.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
d.styles['Normal'].font.name='Aptos';d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6)
for st,sz in [('Title',25),('Heading 1',17),('Heading 2',13),('Heading 3',11.5),('Heading 4',10.5),('Heading 5',10.5)]:d.styles[st].font.name='Aptos Display';d.styles[st].font.size=Pt(sz);d.styles[st].font.color.rgb=RGBColor(0,0,0)
title_style=d.styles['Title'];title_style.font.color.rgb=RGBColor(0,0,0)
title_ppr=title_style._element.pPr
if title_ppr is not None:
 title_border=title_ppr.find(qn('w:pBdr'))
 if title_border is not None:title_ppr.remove(title_border)
p=d.add_paragraph();p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
title_run=p.add_run('B2B Wholesale Management System');title_run.font.color.rgb=RGBColor(0,0,0)
p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Assignment 2 · QT2 · Semester 1 2026–2027');r.italic=True
d.add_paragraph('Database, desktop administration, Razor Pages portal, transaction integrity, tests, and submission evidence.')
add_md(d)
for appendix in ('verification.md','testing.md','demo-checklist.md'):
 d.add_page_break()
 d.add_heading(appendix.replace('.md','').replace('-',' ').title(),1)
 for line in (ROOT/'docs'/appendix).read_text(encoding='utf-8').splitlines():
  s=line.strip()
  if not s:continue
  if s.startswith('!['):continue
  if s.startswith('# '):d.add_heading(s[2:],2)
  elif s.startswith('## '):d.add_heading(s[3:],3)
  elif s.startswith('- '):d.add_paragraph(s[2:],style='List Bullet')
  else:d.add_paragraph(s.replace('**','').replace('`',''))
d.add_page_break();d.add_heading('Verified application screenshots',1)
d.add_heading('Agent Portal login page',2)
d.add_picture(str(ROOT/'docs'/'screenshots'/'portal-login-final.png'),width=Inches(6.15))
d.add_paragraph('Actual final Agent sign-in screen with seven-color NORTHSTAR wordmark, 3D product artwork and a 30-second optional animation.')
d.add_page_break()
d.add_heading('Back-office Admin login form',2)
d.add_picture(str(ROOT/'docs'/'screenshots'/'admin-login-v4.png'),width=Inches(3.8))
d.add_paragraph('Actual compiled WinForms Admin login form. The seeded Admin account authenticated successfully; the SQL-backed inventory and dashboard screens are shown below.')
d.add_heading('Unit and SQL integration test runs',2)
d.add_paragraph('The final MSTest run passed all 26 unit tests. The SQL integration harness passed all 46 checks with an isolated two-product browser checkout and restored the starting table row counts. Current logs are saved in docs/evidence/unit-results.txt and docs/evidence/integration-results.txt.')
for title,file,caption,width in [
 ('Agent catalog with seeded SQL products','portal-catalog.png','Actual authenticated Agent catalog with 15 seeded products and quantity inputs for a multi-item order.',5.55),
 ('Two-line order invoice','portal-invoice.png','Actual order 16 after one transaction: two details totaling 1,045,000 VND.',6.15),
 ('Paid classroom invoice on desktop','payment-confirmed-desktop.png','Actual order 23 after simulated confirmation: two product lines, Paid timestamp, unique reference and printable invoice.',6.15),
 ('Printable paid invoice with registered contact','portal-paid-invoice-final.png','Actual order 55 after simulated confirmation: registered address and phone, Paid timestamp, unique reference and printable invoice.',6.15),
 ('Admin item inventory','admin-dashboard.png','Actual Admin login and SQL-backed item grid after the two-item order.',6.15),
 ('Incoming orders and status','admin-incoming-orders.png','Actual Admin order grid showing order 16 as Pending.',6.15),
 ('Sales and low-stock dashboard','admin-sales-dashboard.png','Actual SQL-backed best sellers, Cashew Nuts low-stock warning, and agent revenue.',6.15),
]:
 d.add_page_break();d.add_heading(title,1)
 d.add_picture(str(ROOT/'docs'/'screenshots'/file),width=Inches(width))
 d.add_paragraph(caption)
d.add_page_break();d.add_heading('Database and system diagrams',0)
for index,(caption,file) in enumerate([('Database entity relationship diagram','erd.png'),('System architecture and deployment boundary','architecture.png')]):
 if index:d.add_page_break()
 d.add_heading(caption,1);d.add_picture(str(ROOT/'docs'/'diagrams'/file),width=Inches(6.3))
 p=d.add_paragraph(f'Figure {index+1}. {caption}.');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if file=='erd.png':d.add_paragraph('Order to OrderDetail and Item to OrderDetail are one-to-many relationships. OrderDetail cascades when its Order is deleted; referenced agents and items remain protected.')
for s in d.sections:
 p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Software Engineering · Assignment 2  |  Page ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,115,130)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
d.save(ROOT/'Report.docx');print(ROOT/'Report.docx')
