from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

ROOT=Path(__file__).parent

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill);tcPr.append(shd)
def set_cell(cell,text,bold=False,color=None):
    cell.text='';p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);r=p.add_run(text);r.bold=bold;r.font.size=Pt(9)
    if color:r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def border_table(table):
    tblPr=table._tbl.tblPr;borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    tblPr.append(borders)
def add_md(doc,path):
    lines=path.read_text(encoding='utf-8').splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('!['):i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(set(v)<=set('-: ') for v in vals): rows.append(vals)
                i+=1
            if not rows:continue
            tbl=doc.add_table(rows=0,cols=len(rows[0]));tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=True
            for ri,row in enumerate(rows):
                cells=tbl.add_row().cells
                for ci,value in enumerate(row):set_cell(cells[ci],value,bold=ri==0,color=(255,255,255) if ri==0 else None);shade(cells[ci],'17365D' if ri==0 else ('F2F6FA' if ri%2==0 else 'FFFFFF'))
            border_table(tbl);doc.add_paragraph().paragraph_format.space_after=Pt(2);continue
        if s.startswith('# '):
            if path.name!='report.md':doc.add_heading(s[2:],1)
        elif s.startswith('## '):
            heading=doc.add_heading(s[3:],2)
            if path.name=='report.md' and s[3:]=='Submission and contribution evidence':heading.paragraph_format.page_break_before=True
        elif s.startswith('### '):doc.add_heading(s[4:],3)
        elif s.startswith('#### '):doc.add_heading(s[5:],4)
        elif s.startswith('##### '):doc.add_heading(s[6:],5)
        elif s.startswith('```'):
            i+=1;code=[]
            while i<len(lines) and not lines[i].strip().startswith('```'):code.append(lines[i]);i+=1
            p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.25);p.paragraph_format.space_after=Pt(8);r=p.add_run('\n'.join(code));r.font.name='Consolas';r.font.size=Pt(8);r.font.color.rgb=RGBColor(45,65,82)
        elif s.startswith('- '):doc.add_paragraph(s[2:],style='List Bullet')
        elif len(s)>2 and s[0].isdigit() and s[1:3]=='. ':doc.add_paragraph(s[3:],style='List Number')
        else:doc.add_paragraph(s.replace('**','').replace('`',''))
        i+=1
def build(name,source,subtitle):
    d=Document();sec=d.sections[0];sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.7);sec.left_margin=Inches(.78);sec.right_margin=Inches(.78)
    styles=d.styles;styles['Normal'].font.name='Aptos';styles['Normal'].font.size=Pt(10.5);styles['Normal'].paragraph_format.space_after=Pt(6)
    for st,size in [('Title',25),('Heading 1',17),('Heading 2',13),('Heading 3',11.5),('Heading 4',10.5),('Heading 5',10.5)]:
        styles[st].font.name='Aptos Display';styles[st].font.size=Pt(size);styles[st].font.color.rgb=RGBColor(0,0,0)
    title_style=styles['Title'];title_style.font.color.rgb=RGBColor(0,0,0)
    title_ppr=title_style._element.pPr
    if title_ppr is not None:
        title_border=title_ppr.find(qn('w:pBdr'))
        if title_border is not None:title_ppr.remove(title_border)
    p=d.add_paragraph();p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title_run=p.add_run('Smart Ticketing System Assignment 1');title_run.font.color.rgb=RGBColor(0,0,0)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run(subtitle);r.italic=True;r.font.color.rgb=RGBColor(89,105,122)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run('QT2 · Semester 1 2026–2027').bold=True
    d.add_paragraph('Requirements, UML, architecture, implemented demo, verification, and honest submission status.')
    add_md(d,source)
    for appendix in ('verification.md','requirements.md','use-case-descriptions.md','copilot-prompts.md','visual-studio-demo.md'):
        if appendix != 'requirements.md':d.add_page_break()
        add_md(d,ROOT/'docs'/appendix)
    for caption,file in [('TVM start screen connected to SQL Server','tvm-start-v6.png'),('Scannable wallet demo QR before payment','tvm-wallet-qr-v2.png'),('Final payment confirmation and saved ticket','payment-confirmed-demo-v2.png')]:
        d.add_page_break();d.add_heading(caption,1)
        d.add_picture(str(ROOT/'docs'/'screenshots'/file),width=Inches(5.75))
        if file=='payment-confirmed-demo-v2.png':d.add_paragraph('Actual SQL-connected WinForms result: the confirmed Wallet ticket is saved as Paid with a valid expiry; the next-journey action is available.')
        else:d.add_paragraph('Actual application screenshot captured from the compiled WinForms program.')
    d.add_page_break();d.add_heading('Design diagrams and UI wireframe',0)
    diagrams=[('Use case diagram','use-case.png'),('Passenger purchase activity','passenger-activity.png'),('Payment gateway activity','gateway-activity.png'),('Purchase sequence diagram','sequence.png'),('Transaction state chart','state.png'),('Domain class diagram','class.png'),('System architecture','architecture.png'),('Deployment diagram','deployment.png'),('TVM touchscreen wireframe','ui-wireframe.png')]
    for i,(caption,file) in enumerate(diagrams):
        if i:d.add_page_break()
        d.add_heading(caption,1)
        if file=='passenger-activity.png':d.add_picture(str(ROOT/'docs'/'diagrams'/file),height=Inches(8.1))
        elif file=='state.png':d.add_picture(str(ROOT/'docs'/'diagrams'/file),height=Inches(8.0))
        else:d.add_picture(str(ROOT/'docs'/'diagrams'/file),width=Inches(6.35))
        p=d.add_paragraph(f'Figure {i+1}. {caption}.');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for sec in d.sections:
        f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=f.add_run('Software Engineering · Assignment 1  |  Page ');r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,115,130)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    target=ROOT/name;d.save(target);print(target)

build('Report.docx',ROOT/'docs/report.md','Smart Ticketing System | Assignment 1')
