"""Rebuild workbooks and Word/PDF reports. Requires openpyxl, python-docx, reportlab."""
from pathlib import Path
import re, json, html, shutil, xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image as PILImage
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[1]
NAME='Phạm Trung Nghĩa'; SID='524H0172'; PROFILE='https://github.com/Nghiaktvn'; DATE='23 October 2026'
old=load_workbook(ROOT/'Assignment1/TVM_Clarification_and_Requirements.xlsx',data_only=True)
QA=[list(r) for r in old['Clarification Q&A'].iter_rows(min_row=5,max_row=16,values_only=True)]
REQ=[list(r) for r in old['Requirements'].iter_rows(min_row=4,max_row=19,values_only=True)]
QMAP=['FR-01, DR-02','FR-02, NFR-06','FR-03, DR-03, DR-05','FR-04, DR-04','FR-07, DR-07','NFR-07','NFR-06','NFR-01..05','FR-02','FR-07, FR-08','FR-02, FR-03','FR-06, NFR-06']
RMAP=['Q1/Q8 → UC-01/06','Q2/Q9 → UC-02/03/07','Q3/Q7 → UC-02/03/05','Q4/Q10 → UC-08','Q4/Q7 → UC-09','Q7/Q12 → UC-04/05','Q3/Q5/Q10 → UC-05','Q10/Q11 → UC-04','Q8 → UC-01','Q8 → UC-06','Q1/Q8 → UC-06','Q3/Q8 → UC-02/03','Q10 → UC-08','Q2/Q7/Q12 → UC-02/03','Q6 → UC-01','Q11 → UC-04']
DR=[['DR-01','At most one ticket per confirmed payment.','Duplicate callbacks and ticket issue are idempotent.'],['DR-02','The active route fare is captured once.','Immutable FareSnapshot.'],['DR-03','Payment references and barcodes are unique.','SQL UNIQUE constraints and same-reference retry.'],['DR-04','Ticket QR uses an opaque ID and service-day validity.','No personal data; expiry displayed.'],['DR-05','Unknown payments reconcile before a new attempt.','Unknown state cannot be silently cancelled.'],['DR-06','Payment QR binds transaction, amount and method and expires.','Demo QR expires after three minutes.'],['DR-07','Printer failure requires compensation.','Same-ticket reprint or recorded simulated refund.']]
UC=[]
for line in (ROOT/'Assignment1/docs/use-case-descriptions.md').read_text(encoding='utf-8-sig').splitlines():
    if line.startswith('| UC-'): UC.append([x.strip() for x in line.strip('|').split('|')])
A1=[('1','Clarification Q&A',5,'2.1','Artifact verified'),('2','Requirements specification',15,'2.2','Artifact verified'),('3','Use cases',10,'3.1','Artifact reviewed'),('4','Activity diagrams',10,'3.2','Artifact reviewed'),('5','Detailed system modeling',15,'3.3','Current class/state/sequence reviewed'),('6','UI/UX',5,'4.1','Built and visually checked'),('7','Architecture',10,'4.2','Demo/production boundaries stated'),('8','Visual Studio engineering',10,'5.1','Generated artifacts; actual VS captures pending'),('9','GitHub and AI tooling',10,'5.2','AI log supplied; genuine group GitHub history pending'),('10','C# and MSSQL demo',10,'5.3','21 domain tests; SQL-connected demo')]
A2=[('DB','Database and integrity',10,'3','Schema/seed and SQL verified'),('1.1','Admin login',5,'4.1','Role and lockout SQL checks'),('1.2','Item, Agent and Order management',10,'4.2','CRUD, row version and stock SQL checks'),('1.3','Shipping manifest',5,'4.3','HTML export and print stylesheet verified'),('1.4','Dashboard and filters',5,'4.4','SQL agent/date filters verified'),('3.1','Agent login',5,'5.2','Browser and SQL checks'),('3.2','Master-detail transaction',15,'5.3','Rollback/replay/concurrency verified'),('3.3','My Orders and invoice',10,'5.4','Browser and ownership checks'),('3.4','Responsive UI',5,'5.5','Desktop and 304px view checked'),('5.1','Meaningful unit tests',15,'6','26/26 passed'),('6a','Configuration',5,'7','External config files used'),('6b','Git collaboration',5,'8','Actual repository and PR history pending'),('6c','Professional report',5,'1-10','Generated; full team record pending')]
def tests(n):
    folder=ROOT/f'Assignment{n}/tests'
    files=list(folder.glob('**/TestResults/*.trx'))
    if not files:return []
    return [(x.attrib['testName'],x.attrib['outcome']) for x in ET.parse(files[0]).findall('.//{http://microsoft.com/schemas/VisualStudio/TeamTest/2010}UnitTestResult')]
def table(headers,rows):
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join('' if x is None else str(x).replace('|','/') for x in row)+' |' for row in rows])+'\n'
def team_table():
    return table(['Member','Student ID','Actual responsibilities','Contribution'],[[NAME,SID,'Confirmed submission contact; feature ownership pending group confirmation','Pending'],['Remaining 2–4 members','Not supplied','Group must provide real names and ownership','Pending']])
def sheet(wb,name,headers,rows,widths=None,editable=(),lecturer=()):
    ws=wb.create_sheet(name);ws.append([name+' | '+DATE]);ws.append(['Verified artifacts and pending team/lecturer evidence are recorded separately.']);ws.append(headers)
    for row in rows:ws.append(row)
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(headers));ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=len(headers))
    ws.row_dimensions[1].height=30;ws.row_dimensions[2].height=30;ws.freeze_panes='A4';ws.sheet_view.showGridLines=False
    for row in ws:
        for c in row:
            c.alignment=Alignment(wrap_text=True,vertical='top');c.font=Font(name='Calibri',size=10,color='1F4E78' if c.column in editable or c.column in lecturer else '202C2C')
            if c.row==1:c.fill=PatternFill('solid',fgColor='183638');c.font=Font(name='Calibri',size=16,bold=True,color='FFFFFF')
            elif c.row==3:c.fill=PatternFill('solid',fgColor='24715D');c.font=Font(name='Calibri',bold=True,color='FFFFFF')
            elif c.row>3:
                c.fill=PatternFill('solid',fgColor='E2F0D9' if c.column in lecturer else 'FFF2CC' if c.column in editable else 'F3F7F2' if c.row%2==0 else 'FFFFFF')
    from openpyxl.utils import get_column_letter
    for i in range(len(headers)):ws.column_dimensions[get_column_letter(i+1)].width=(widths or [27]*len(headers))[i]
    import math
    for r in range(4,ws.max_row+1):
        wrapped=max((math.ceil(len(str(c.value or ''))/max(12,ws.column_dimensions[c.column_letter].width-2)) for c in ws[r]),default=1)
        ws.row_dimensions[r].height=min(170,max(42,wrapped*13+12))
    if rows:ws.auto_filter.ref=f'A3:{ws.cell(ws.max_row,len(headers)).coordinate}'
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.print_title_rows='1:3'
    return ws
def workbook(n):
    wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    team=sheet(wb,'Team',['Name','Student ID','Role / actual ownership','Contribution %','Commit count','GitHub profile'],[[NAME,SID,'Confirmed contact; actual feature ownership pending',None,'=COUNTIF(\'Commit Log\'!C4:C500,B4)',PROFILE]]+[[None,None,None,None,None,None] for _ in range(4)],[26,17,57,18,18,40],editable=(1,2,3,4,6))
    for r in range(5,9):team.cell(r,5,f'=IF(B{r}="","",COUNTIF(\'Commit Log\'!C4:C500,B{r}))')
    team.append(['TOTAL',None,'3–5 actual members required','=SUM(D4:D8)',None,None]);team.append(['Group number',None,'Pending group confirmation']);team.append(['Web option' if n==2 else 'Assignment','Option B — Razor Pages' if n==2 else 'Smart Ticketing'])
    dv=DataValidation(type='decimal',operator='between',formula1=0,formula2=100);dv.showErrorMessage=True;team.add_data_validation(dv);dv.add('D4:D8')
    criteria=A1 if n==1 else A2
    if n==1:
        sheet(wb,'Clarification Q&A',['ID','Question','Assumed answer','Design impact','Requirement IDs'],[r+[QMAP[i]] for i,r in enumerate(QA)],[10,45,65,55,28],editable=(2,3,4,5))
        sheet(wb,'Requirements',['ID','Type','Requirement','Acceptance target','Priority','Source / use case'],[r+['Must',RMAP[i]] for i,r in enumerate(REQ)]+[[r[0],'Domain',r[1],r[2],'Must','Q3/Q4/Q10 / UC-01'] for r in DR],[15,20,75,72,15,32],editable=(3,4,5,6))
        sheet(wb,'Use Cases',['Use case','Preconditions','Main flow','Alternate / exception','Postcondition','Owner'],[r+[None] for r in UC],[30,45,70,70,48,28],editable=(6,))
        sheet(wb,'Artifact Tracker',['Artifact','Path','Tool','Status','Owner','Completion %'],[['Q&A/requirements','A1_Team_Workbook.xlsx','Excel','Verified',None,100],['Use case/activity diagrams','docs/diagrams','DOT + PNG','Reviewed',None,100],['Sequence/state/class','docs/diagrams','DOT / Python','Reviewed',None,100],['UI and architecture','docs/screenshots; diagrams','WinForms','Built and checked',None,100],['Engineering','engineering','JSON / C# / DGML','Actual VS UI capture pending',None,75],['Git collaboration','Real group repository','GitHub','Pending',None,0],['Demo and tests','SmartTicketing.sln','C# / SQL / MSTest','21 tests passed',None,100],['Report','Report.docx','Word','Generated; group details pending',None,90]],[30,55,25,48,28,18],editable=(5,6))
        sheet(wb,'Consistency Check',['Cross-check','Evidence','Result','Reviewer','Lecturer feedback'],[['Requirements → use cases','FR-01..07 / UC-01..09','Reviewed',None,None],['Class/state/sequence → code','TicketTransaction and TicketStore','Reviewed',None,None],['All payment methods','Card / BankQr / Wallet','Reviewed',None,None],['Unknown state never issues ticket','TicketTransactionTests','Passed',None,None],['Demo and production boundary','Report 4.2; architecture.png','Explicit',None,None],['3–5 members and 100% contribution','Team sheet','Pending',None,None]],[42,65,22,28,45],editable=(4,),lecturer=(5,))
        sheet(wb,'AI Prompt Log',['Date','Tool','Actual prompt summary','Kept / changed','Validation'],[[DATE,'Codex','Complete PDFs and rubric','Domain, tests, workbooks and reports','Build / MSTest / SQL demo'],[DATE,'Codex','Upgrade UI/UX and functional demos','Kiosk states, original design and 30-second media','Native screenshots / browser'],[DATE,'image_gen','Joyful MetroGo 3D transit scene','Original artwork; timed motion graphics','Image inspection / metadata']],[24,20,65,65,55])
    else:
        sheet(wb,'Tasks & Division',['Task','Feature','Marks','Responsible student','Branch / PR','Status','Workload share %'],[[a,b,c,None,None,e,None] for a,b,c,d,e in criteria],[12,52,12,32,50,64,22],editable=(4,5,7))
        sheet(wb,'Database',['Table','PK','Keys and constraints','Seed rows','Evidence'],[['Item','ItemID','Positive price; nonnegative stock; RowVersion',15,'Schema + SQL harness'],['Agent','AgentID','Unique phone; soft deactivation',15,'Registration rollback tested'],['Order','OrderID','Agent FK; status check; unique RequestKey when set',15,'Atomic/replayed orders tested'],['OrderDetail','ID','Order and Item FKs; unique order/item; positive qty/price',15,'Mid-transaction rollback tested'],['Users','UserID','Unique username/email; role/agent CHECK; PBKDF2',16,'Role and lockout tests']],[20,18,82,18,65])
        data=[['Enough stock','2 × 100 + 3 × 75','425 total; both stocks deducted','Passed'],['Insufficient stock','Second quantity 999','No header/details; stocks unchanged','Passed'],['Mid-transaction failure','Controlled second-detail trigger failure','Undo header, first detail and first deduction','Passed'],['Replay same request key','Duplicate checkout','Same order; no second deduction','Passed'],['Concurrent last unit','Two simultaneous orders','Exactly one succeeds; stock zero','Passed'],['Cancel then repeat cancel','Status compare/update','Stock returned once only','Passed'],['Foreign invoice','Agent B requests A invoice','No access','Passed'],['Actual browser checkout','Order #23: 2 Arabica + 3 Robusta','1,045,000 VND','Passed']]
        sheet(wb,'Transaction Tests',['Scenario','Input','Expected / actual','Result'],data,[33,65,80,18])
        sheet(wb,'Unit Tests',['Test name','Business logic','Expected','Actual'],[[a,'OrderRules or PasswordHash; inspect test method','Assertion passes',b] for a,b in tests(2)],[62,68,30,20])
        sheet(wb,'Branches & PRs',['Feature','Suggested branch','Actual PR','Reviewer','Status'],[[a,b,None,None,'Suggested workflow; no historical activity invented'] for a,b in [('Database','feature/database-integrity'),('Desktop','feature/admin-workspace'),('Portal','feature/agent-portal'),('Tests and documents','feature/testing-docs')]],[32,48,65,28,70],editable=(2,3,4))
        sheet(wb,'Config Check',['Check','Location','How read / verified','Result','Lecturer feedback'],[['Admin connection','src/Wholesale.AdminWinForms/App.config','XML configured connection','Verified',None],['Portal connection','src/AgentPortal/appsettings.json','GetConnectionString; environment overrides','Verified',None],['No application hardcoded connection','*.cs files','Connections read from configuration','Verified source',None],['Passwords','Users.PasswordHash','Shared PBKDF2 + bounded format + constant-time match','Unit/SQL tests',None],['Key material','App_Data/keys','Ignored; excluded from packages','Verified',None]],[32,70,78,24,35],lecturer=(5,))
        sheet(wb,'AI Usage Log',['Date','Tool','Actual prompt summary','Kept / changed','Validation'],[[DATE,'Codex','Complete both PDF assignments','Code and rubric artifacts','26 unit + 33 SQL checks'],[DATE,'Codex','MilkyBloom-like layout; preserve features','Original home/cards/3D media','Desktop + narrow browser demo'],[DATE,'image_gen','Joyful 3D wholesale delivery scene','Original artwork + 30-second motion graphics','Inspected image and duration']],[24,20,65,65,55])
    sheet(wb,'Commit Log',['Date / lab','Commit SHA','Student ID','Branch','PR URL','Evidence'],[['No entries recorded',None,None,None,None,'Local main has no commits or remote. Add genuine group evidence.']],[25,48,20,40,65,80],editable=(1,2,3,4,5,6))
    sheet(wb,'Progress Log',['Date / lab','Plan','Actual work','Evidence','Student','Lecturer feedback'],[[DATE,'Close rubric gaps and upgrade demos','Built UI, domain rules, tests and documents','docs/evidence and screenshots',None,None]]+[[x,'Record actual team work',None,None,None,None] for x in ['Lab 7 / Week 1','Lab 8 / Week 2','Lab 9 / Week 3','Lab 10 / Week 4']],[26,65,70,60,28,45],editable=(2,3,4,5),lecturer=(6,))
    mark=sheet(wb,'Marking',['Task','Criterion','Maximum','Self-mark','Lecturer mark','Status / evidence','Section'],[[a,b,c,None,None,e,d] for a,b,c,d,e in criteria],[12,50,15,20,22,75,20],editable=(4,),lecturer=(5,))
    end=3+len(criteria);mark.append(['TOTAL',None,f'=SUM(C4:C{end})',f'=IF(COUNT(D4:D{end})=0,"",SUM(D4:D{end}))',f'=IF(COUNT(E4:E{end})=0,"",SUM(E4:E{end}))','Marks are not guaranteed'])
    sheet(wb,'Summary',['Metric','Value','Meaning'],[['Maximum score',100,'Rubric total'],['Self-mark',f'=IF(COUNT(Marking!D4:D{end})=0,"",SUM(Marking!D4:D{end}))','Team review required'],['Lecturer mark',f'=IF(COUNT(Marking!E4:E{end})=0,"",SUM(Marking!E4:E{end}))','Lecturer input only'],['QT2 weight',.2 if n==1 else .3,'Assignment fraction'],['Weighted QT2 percentage points','=IF(B6="","",B6*B7)','Lecturer score times weight'],['Member count','=COUNTA(Team!B4:B8)','Must be 3–5'],['Contribution total','=SUM(Team!D4:D8)','Must be 100'],['Team readiness','=IF(AND(B9>=3,B9<=5,B10=100),"Ready","Missing team details")','Do not fabricate members or contributions'],['GitHub repository','Pending','Provided URL is an account profile'],['Manual evidence','Actual Visual Studio screenshots' if n==1 else 'Genuine group repository and PRs','See report']],[44,53,90])
    target=ROOT/f'Assignment{n}/A{n}_Team_Workbook.xlsx';wb.save(target);print(target)

def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(set(x)<=set('-: ') for x in vals):rows.append(vals)
                i+=1
            yield 'table',rows;continue
        if s.startswith('```'):
            data=[];i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):data.append(lines[i]);i+=1
            yield 'code','\n'.join(data)
        elif s.startswith('#'):
            m=re.match(r'^(#+) (.*)',s);yield 'heading',(len(m[1]),m[2])
        elif s.startswith('!['):
            m=re.match(r'!\[(.*?)\]\((.*?)\)',s);yield 'image',(m[1],m[2])
        else:yield 'text',s
        i+=1
def reports(n,text):
    folder=ROOT/f'Assignment{n}';(folder/'docs/report.md').write_text(text,encoding='utf-8')
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.68);sec.left_margin=sec.right_margin=Inches(.72)
    d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
    for level,size in [(1,18),(2,14),(3,12),(4,11)]:d.styles[f'Heading {level}'].font.size=Pt(size);d.styles[f'Heading {level}'].font.color.rgb=RGBColor.from_string('183638')
    d.add_heading('SMART TICKETING' if n==1 else 'NORTHSTAR WHOLESALE',0);d.add_paragraph(f'Software Engineering Lab · Assignment {n} · Semester 1 2026–2027\n{NAME} · {SID}\n{DATE}\n{PROFILE}')
    d.add_paragraph('Verified work and pending group/lecturer evidence are explicitly separated. The full 3–5 member roster, contributions and genuine GitHub collaboration record must be supplied by the group.')
    d.add_page_break();d.add_heading('Table of Contents',1);p=d.add_paragraph();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-3" \\h \\z \\u');p._p.append(fld);d.add_paragraph('In Word, update fields to populate the contents and page numbers. All numbered headings are real Word headings.');d.add_page_break()
    for kind,val in blocks(text):
        if kind=='heading':
            level,title=val
            if level==1:continue
            p=d.add_heading(title,min(level-1,4))
            if level==2 and re.match(r'^\d+\.',title):p.paragraph_format.page_break_before=True
        elif kind in ('text','code'):
            p=d.add_paragraph(val)
            if kind=='code':
                for r in p.runs:r.font.name='Consolas';r.font.size=Pt(8)
        elif kind=='table':
            t=d.add_table(rows=0,cols=len(val[0]));t.style='Light Shading Accent 1'
            for ri,row in enumerate(val):
                cells=t.add_row().cells
                for ci,value in enumerate(row):
                    cells[ci].text=value
                    for p in cells[ci].paragraphs:
                        p.paragraph_format.space_after=Pt(3)
                        for r in p.runs:r.font.size=Pt(8.5);r.bold=ri==0
                if ri==0:t.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            d.add_paragraph()
        elif kind=='image':
            caption,relative=val;p=folder/'docs'/relative
            if p.exists():
                with PILImage.open(p) as im:w,h=im.size
                d.add_picture(str(p),width=Inches(min(6.7,7.8*w/h)));d.add_paragraph(caption)
    update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');d.settings._element.append(update)
    for sec in d.sections:
        p=sec.footer.paragraphs[0];p.add_run(f'Assignment {n} · {SID} | Page ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    d.save(folder/'Report.docx')
    if n==1:shutil.copy2(folder/'Report.docx',folder/'524H0172_PhamTrungNghia.docx')
    pdfmetrics.registerFont(TTFont('LabFont','C:/Windows/Fonts/arial.ttf'));pdfmetrics.registerFont(TTFont('LabBold','C:/Windows/Fonts/arialbd.ttf'))
    normal=ParagraphStyle('Normal',fontName='LabFont',fontSize=9,leading=13,spaceAfter=7,textColor=colors.HexColor('#183638'));small=ParagraphStyle('Cell',parent=normal,fontSize=7,leading=9,spaceAfter=0)
    story=[]
    for kind,val in blocks(text):
        if kind=='heading':
            level,title=val;size=18 if level==1 else 14 if level==2 else 11
            style=ParagraphStyle('H',fontName='LabBold',fontSize=size,leading=size+5,spaceBefore=13,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor('#183638'));story.append(Paragraph(html.escape(title),style))
        elif kind in ('text','code'):story.append(Paragraph(html.escape(val).replace('\n','<br/>'),normal))
        elif kind=='table':
            t=Table([[Paragraph(html.escape(x),small) for x in r] for r in val],colWidths=[505/len(val[0])]*len(val[0]),repeatRows=1)
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDEBE0')),('GRID',(0,0),(-1,-1),.3,colors.HexColor('#D9E3D7')),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6)]));story.extend([t,Spacer(1,10)])
        elif kind=='image':
            caption,relative=val;p=folder/'docs'/relative
            if p.exists():
                with PILImage.open(p) as im:w,h=im.size
                ratio=min(505/w,610/h);story.extend([KeepTogether([Image(str(p),width=w*ratio,height=h*ratio),Paragraph(html.escape(caption),normal)]),Spacer(1,10)])
    def footer(canvas,doc):canvas.setFont('LabFont',7);canvas.drawString(45,23,f'Assignment {n} · {SID}');canvas.drawRightString(550,23,f'Page {doc.page}')
    SimpleDocTemplate(str(folder/'Report.pdf'),pagesize=(595,842),leftMargin=45,rightMargin=45,topMargin=40,bottomMargin=40).build(story,onFirstPage=footer,onLaterPages=footer)
    print(folder/'Report.docx');print(folder/'Report.pdf')

if __name__=='__main__':
    from report_content import a1_content,a2_content
    for n in (1,2):workbook(n)
    reports(1,a1_content());reports(2,a2_content())
